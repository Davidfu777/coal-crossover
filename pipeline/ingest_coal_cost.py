from __future__ import annotations

import re

import pandas as pd
from openpyxl import load_workbook

from .countries import normalize_country
from .paths import IEA_NEA_WORKBOOK, PINK_SHEET, US_GDP_DEFLATOR

MMBTU_PER_MT_AT_6000_KCAL_PER_KG = 6000 * 1000 * 3.968320719 / 1_000_000
PINK_SHEET_START = "2023M01"
PINK_SHEET_END = "2025M12"
COAL_PRICE_PERCENTILES = tuple(range(0, 101))


def usd_2018_to_2025_factor() -> float:
    deflator = pd.read_csv(US_GDP_DEFLATOR)
    united_states = deflator.loc[deflator["country_iso3"].eq("USA")].set_index("year")["value"]
    if 2018 not in united_states or 2025 not in united_states:
        raise ValueError("World Bank U.S. GDP deflator does not cover 2018 and 2025")
    return float(united_states.loc[2025] / united_states.loc[2018])


def _monthly_coal_prices_usd_per_mmbtu() -> pd.DataFrame:
    raw = pd.read_excel(PINK_SHEET, sheet_name="Monthly Prices", header=4)
    date_column = raw.columns[0]
    prices = raw[[date_column, "Coal, Australian", "Coal, South African **"]].copy()
    prices.columns = ["month", "australia_usd_per_mt", "south_africa_usd_per_mt"]
    prices = prices.loc[prices["month"].astype(str).between(PINK_SHEET_START, PINK_SHEET_END)]
    for column in ["australia_usd_per_mt", "south_africa_usd_per_mt"]:
        prices[column] = pd.to_numeric(prices[column], errors="coerce")
    if len(prices) != 36 or prices.iloc[:, 1:].isna().any().any():
        raise ValueError("Pink Sheet coal benchmark window is incomplete")

    prices["blend_usd_per_mt"] = prices[["australia_usd_per_mt", "south_africa_usd_per_mt"]].mean(
        axis=1
    )
    result = pd.DataFrame(index=prices.index)
    for benchmark, column in {
        "australia": "australia_usd_per_mt",
        "south_africa": "south_africa_usd_per_mt",
        "australia_south_africa_blend": "blend_usd_per_mt",
    }.items():
        result[benchmark] = prices[column] / MMBTU_PER_MT_AT_6000_KCAL_PER_KG
    return result


def load_coal_price_quantiles(
    percentiles: tuple[int, ...] = COAL_PRICE_PERCENTILES,
) -> pd.DataFrame:
    prices = _monthly_coal_prices_usd_per_mmbtu()
    rows: list[dict[str, object]] = []
    for benchmark in prices.columns:
        values = prices[benchmark]
        for percentile in percentiles:
            rows.append(
                {
                    "coal_price_benchmark": benchmark,
                    "percentile": int(percentile),
                    "coal_price_usd_per_mmbtu": float(
                        values.quantile(percentile / 100, interpolation="linear")
                    ),
                    "coal_price_months": int(values.count()),
                    "coal_price_window": f"{PINK_SHEET_START}-{PINK_SHEET_END}",
                    "quantile_method": "linear interpolation of the empirical monthly distribution",
                    "coal_price_source": "World Bank Commodity Price Data (Pink Sheet)",
                }
            )
    return pd.DataFrame(rows)


def load_coal_price_benchmarks() -> pd.DataFrame:
    prices = _monthly_coal_prices_usd_per_mmbtu()
    rows: list[dict[str, object]] = []
    for benchmark in prices.columns:
        values = prices[benchmark]
        rows.append(
            {
                "coal_price_benchmark": benchmark,
                "coal_price_p10_usd_per_mmbtu": float(values.quantile(0.10)),
                "coal_price_usd_per_mmbtu": float(values.median()),
                "coal_price_p90_usd_per_mmbtu": float(values.quantile(0.90)),
                "coal_price_months": int(values.count()),
                "coal_price_window": f"{PINK_SHEET_START}-{PINK_SHEET_END}",
                "coal_price_source": "World Bank Commodity Price Data (Pink Sheet)",
            }
        )
    return pd.DataFrame(rows)


def _normalize_technology(value: object) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", text)


def load_iea_nea_coal_om() -> pd.DataFrame:
    workbook = load_workbook(IEA_NEA_WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["3.12"]
    inflation_factor = usd_2018_to_2025_factor()
    records: list[dict[str, object]] = []
    country: str | None = None
    in_non_oecd_section = False
    for row in sheet.iter_rows(min_row=6, values_only=True):
        raw_country = row[1]
        technology = row[2]
        om = row[13]
        if raw_country == "Non-OECD countries":
            in_non_oecd_section = True
            country = None
            continue
        if isinstance(raw_country, str) and raw_country.strip():
            country = raw_country.strip()
        if not isinstance(technology, str) or not isinstance(om, (int, float)):
            continue
        normalized_technology = _normalize_technology(technology)
        records.append(
            {
                "iea_country": country,
                "country_key": normalize_country(country),
                "iea_non_oecd_section": in_non_oecd_section,
                "iea_technology": technology.strip(),
                "iea_technology_key": normalized_technology,
                "contains_ccus": "ccus" in normalized_technology,
                "contains_lignite": "lignite" in normalized_technology,
                "coal_om_reported_2018_usd_per_mwh": float(om),
                "usd_2018_to_2025_factor": inflation_factor,
                "coal_om_usd_per_mwh": float(om) * inflation_factor,
                "coal_om_source": "IEA/NEA Projected Costs of Generating Electricity 2020, Table 3.12",
                "coal_om_cost_year": 2025,
            }
        )
    result = pd.DataFrame(records)
    if result.empty:
        raise ValueError("No coal O&M observations were parsed from IEA/NEA Table 3.12")
    return result


def _unit_uses_ccs(row: pd.Series) -> bool:
    return "ccs" in f"{row.get('combustion_technology', '')} {row.get('coal_type', '')}".lower()


def _technology_score(unit: pd.Series, candidate: pd.Series) -> int:
    gem = str(unit.get("combustion_technology", "")).lower()
    coal_type = str(unit.get("coal_type", "")).lower()
    iea = str(candidate["iea_technology_key"])
    score = 0
    if "ultra-supercritical" in gem and "ultra-supercritical" in iea:
        score += 8
    elif "supercritical" in gem and "supercritical" in iea and "ultra" not in iea:
        score += 8
    elif "subcritical" in gem and ("subcritical" in iea or iea == "pulverised"):
        score += 8
    elif "igcc" in gem and "igcc" in iea:
        score += 8
    elif "unknown" in gem or not gem:
        score += 1
    if "lignite" in coal_type and candidate["contains_lignite"]:
        score += 4
    elif "lignite" not in coal_type and not candidate["contains_lignite"]:
        score += 2
    return score


def _select_om_row(unit: pd.Series, observations: pd.DataFrame) -> tuple[pd.Series, str]:
    ccs = _unit_uses_ccs(unit)
    compatible = observations.loc[observations["contains_ccus"].eq(ccs)].copy()
    if compatible.empty:
        compatible = observations.copy()

    exact_country = compatible.loc[compatible["country_key"].eq(unit["country_key"])]
    if not exact_country.empty:
        exact_country = exact_country.copy()
        exact_country["score"] = exact_country.apply(
            lambda row: _technology_score(unit, row), axis=1
        )
        return exact_country.sort_values(
            ["score", "coal_om_usd_per_mwh"], ascending=[False, True]
        ).iloc[0], "country_technology"

    compatible = compatible.copy()
    compatible["score"] = compatible.apply(lambda row: _technology_score(unit, row), axis=1)
    best_score = compatible["score"].max()
    technology_pool = compatible.loc[compatible["score"].eq(best_score)]
    if best_score > 1 and not technology_pool.empty:
        selected = technology_pool.iloc[
            (
                technology_pool["coal_om_usd_per_mwh"]
                - technology_pool["coal_om_usd_per_mwh"].median()
            )
            .abs()
            .argmin()
        ]
        selected = selected.copy()
        selected["coal_om_usd_per_mwh"] = float(technology_pool["coal_om_usd_per_mwh"].median())
        selected["iea_country"] = None
        return selected, "technology_median"

    selected = compatible.iloc[0].copy()
    selected["coal_om_usd_per_mwh"] = float(compatible["coal_om_usd_per_mwh"].median())
    selected["iea_country"] = None
    selected["iea_technology"] = "All compatible coal observations"
    return selected, "global_coal_median"


def _coal_price_mapping(unit: pd.Series) -> tuple[str, str]:
    source = str(unit.get("coal_source", "")).lower()
    if "australia" in source:
        return "australia", "explicit_coal_source"
    if "south africa" in source or "richards bay" in source:
        return "south_africa", "explicit_coal_source"
    region = str(unit.get("region", "")).lower()
    if region in {"asia", "oceania"}:
        return "australia", "regional_international_proxy"
    if region in {"europe", "africa"}:
        return "south_africa", "regional_international_proxy"
    return "australia_south_africa_blend", "regional_international_proxy"


def assign_coal_cost_inputs(units: pd.DataFrame) -> pd.DataFrame:
    benchmarks = load_coal_price_benchmarks().set_index("coal_price_benchmark")
    om_observations = load_iea_nea_coal_om()
    records: list[dict[str, object]] = []
    for _, unit in units.iterrows():
        benchmark_name, price_level = _coal_price_mapping(unit)
        benchmark = benchmarks.loc[benchmark_name]
        om_row, om_level = _select_om_row(unit, om_observations)
        heat_rate_mmbtu_per_mwh = float(unit["heat_rate_btu_per_kwh"]) / 1000
        central_fuel = heat_rate_mmbtu_per_mwh * float(benchmark["coal_price_usd_per_mmbtu"])
        low_fuel = heat_rate_mmbtu_per_mwh * float(benchmark["coal_price_p10_usd_per_mmbtu"])
        high_fuel = heat_rate_mmbtu_per_mwh * float(benchmark["coal_price_p90_usd_per_mmbtu"])
        om = float(om_row["coal_om_usd_per_mwh"])
        records.append(
            {
                "unit_id": unit["unit_id"],
                "coal_price_benchmark": benchmark_name,
                "coal_price_source_level": price_level,
                "coal_price_usd_per_mmbtu": float(benchmark["coal_price_usd_per_mmbtu"]),
                "coal_price_p10_usd_per_mmbtu": float(benchmark["coal_price_p10_usd_per_mmbtu"]),
                "coal_price_p90_usd_per_mmbtu": float(benchmark["coal_price_p90_usd_per_mmbtu"]),
                "coal_price_window": benchmark["coal_price_window"],
                "heat_rate_mmbtu_per_mwh": heat_rate_mmbtu_per_mwh,
                "fuel_cost_usd_per_mwh": central_fuel,
                "fuel_cost_p10_usd_per_mwh": low_fuel,
                "fuel_cost_p90_usd_per_mwh": high_fuel,
                "coal_om_usd_per_mwh": om,
                "coal_om_source_level": om_level,
                "coal_om_country_used": om_row["iea_country"],
                "coal_om_technology_used": om_row["iea_technology"],
                "coal_operating_cost_usd_per_mwh": central_fuel + om,
                "coal_operating_cost_p10_usd_per_mwh": low_fuel + om,
                "coal_operating_cost_p90_usd_per_mwh": high_fuel + om,
                "coal_cost_basis": "fuel plus all-in O&M; excludes sunk capital, carbon price, and sustaining capital",
            }
        )
    return pd.DataFrame(records)
