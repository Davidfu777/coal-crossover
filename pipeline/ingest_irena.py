from __future__ import annotations

from dataclasses import dataclass

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .countries import irena_cost_region, is_oecd, normalize_country
from .paths import IRENA_WORKBOOK


@dataclass(frozen=True)
class IrenaInputs:
    capex: pd.DataFrame
    wacc: pd.DataFrame
    wind_om: pd.DataFrame
    solar_om_oecd: float
    solar_om_non_oecd: float
    lifetime_years: int


def _number(value: object) -> float:
    if value is None:
        return np.nan
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(" ", "").replace(",", "")
    try:
        return float(text)
    except ValueError:
        return np.nan


def _capex_table(workbook, sheet: str, technology: str) -> pd.DataFrame:
    ws = workbook[sheet]
    rows: list[dict[str, object]] = []
    for row in ws.iter_rows(values_only=True):
        label = row[1] if len(row) > 1 else None
        p05 = _number(row[5] if len(row) > 5 else None)
        mean = _number(row[6] if len(row) > 6 else None)
        p95 = _number(row[7] if len(row) > 7 else None)
        if isinstance(label, str) and np.isfinite(mean):
            rows.append(
                {
                    "technology": technology,
                    "cost_geography": label.strip(),
                    "capex_p05_usd_per_kw": p05,
                    "capex_usd_per_kw": mean,
                    "capex_p95_usd_per_kw": p95,
                    "capex_source_level": "country_or_region",
                }
            )
    return pd.DataFrame(rows)


def _global_capex(workbook, sheet: str, technology: str) -> dict[str, object]:
    ws = workbook[sheet]
    years = [ws.cell(5, column).value for column in range(1, ws.max_column + 1)]
    column = years.index(2025) + 1
    return {
        "technology": technology,
        "cost_geography": "Global",
        "capex_p05_usd_per_kw": float(ws.cell(6, column).value),
        "capex_usd_per_kw": float(ws.cell(7, column).value),
        "capex_p95_usd_per_kw": float(ws.cell(8, column).value),
        "capex_source_level": "global",
    }


def load_irena_inputs() -> IrenaInputs:
    workbook = load_workbook(IRENA_WORKBOOK, read_only=True, data_only=True)
    capex = pd.concat(
        [
            _capex_table(workbook, "Table 2.1", "wind"),
            _capex_table(workbook, "Table 3.1", "solar"),
            pd.DataFrame([_global_capex(workbook, "Fig 2.3", "wind")]),
            pd.DataFrame([_global_capex(workbook, "Fig 3.1", "solar")]),
        ],
        ignore_index=True,
    )

    ws = workbook["Table A2"]
    wacc_rows: list[dict[str, object]] = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        country = row[1] if len(row) > 1 else None
        wind = _number(row[6] if len(row) > 6 else None)
        solar = _number(row[7] if len(row) > 7 else None)
        if isinstance(country, str) and np.isfinite(wind) and np.isfinite(solar):
            wacc_rows.extend(
                [
                    {
                        "country_key": normalize_country(country),
                        "technology": "wind",
                        "wacc_real_after_tax": wind,
                    },
                    {
                        "country_key": normalize_country(country),
                        "technology": "solar",
                        "wacc_real_after_tax": solar,
                    },
                ]
            )
    wacc = pd.DataFrame(wacc_rows)

    ws = workbook["Table A3"]
    wind_om_rows: list[dict[str, object]] = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        country = row[1] if len(row) > 1 else None
        value = _number(row[2] if len(row) > 2 else None)
        if isinstance(country, str) and np.isfinite(value):
            wind_om_rows.append({"om_geography": country.strip(), "wind_om_usd_per_kw_year": value})
    wind_om = pd.DataFrame(wind_om_rows)

    ws = workbook["Table A4"]
    solar_2025 = next(row for row in ws.iter_rows(min_row=4, values_only=True) if row[1] == 2025)

    return IrenaInputs(
        capex=capex,
        wacc=wacc,
        wind_om=wind_om,
        solar_om_oecd=float(solar_2025[2]),
        solar_om_non_oecd=float(solar_2025[3]),
        lifetime_years=25,
    )


def assign_renewable_cost_inputs(plants: pd.DataFrame, inputs: IrenaInputs) -> pd.DataFrame:
    records: list[dict[str, object]] = []
    capex_lookup = inputs.capex.set_index(["technology", "cost_geography"]).to_dict("index")
    wacc_lookup = inputs.wacc.set_index(["technology", "country_key"])[
        "wacc_real_after_tax"
    ].to_dict()
    global_wacc = inputs.wacc.groupby("technology")["wacc_real_after_tax"].median().to_dict()
    wind_om_lookup = inputs.wind_om.set_index("om_geography")["wind_om_usd_per_kw_year"].to_dict()

    for plant in plants.itertuples(index=False):
        for technology in ("solar", "wind"):
            geography = irena_cost_region(plant.country, plant.region, plant.subregion, technology)
            capex = capex_lookup.get((technology, geography))
            if capex is None:
                geography = "Global"
                capex = capex_lookup[(technology, geography)]
            country_key = plant.country_key
            wacc = wacc_lookup.get((technology, country_key), global_wacc[technology])
            wacc_level = "country" if (technology, country_key) in wacc_lookup else "global_median"
            if technology == "solar":
                om = inputs.solar_om_oecd if is_oecd(country_key) else inputs.solar_om_non_oecd
                om_geography = "OECD" if is_oecd(country_key) else "Non-OECD"
            else:
                exact = next(
                    (key for key in wind_om_lookup if normalize_country(key) == country_key), None
                )
                om_geography = exact or ("Other OECD" if is_oecd(country_key) else "Other non-OECD")
                om = wind_om_lookup[om_geography]
            records.append(
                {
                    "plant_id": plant.plant_id,
                    "technology": technology,
                    **capex,
                    "capex_geography_used": geography,
                    "wacc_real_after_tax": float(wacc),
                    "wacc_source_level": wacc_level,
                    "om_usd_per_kw_year": float(om),
                    "om_geography_used": om_geography,
                    "lifetime_years": inputs.lifetime_years,
                    "renewable_cost_source": "IRENA Renewable Power Generation Costs in 2025",
                }
            )
    return pd.DataFrame(records)
